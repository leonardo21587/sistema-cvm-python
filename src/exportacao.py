from __future__ import annotations

from io import BytesIO
from typing import Any
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.indicadores_financeiros import INDICADORES_FINANCEIROS
from src.layouts_cvm import LAYOUT_FINANCEIRA


# ============================================================
# SISTEMA CVM — EXPORTAÇÃO EXCEL
# ============================================================
#
# Este módulo NÃO recalcula demonstrações nem indicadores.
# Ele apenas exporta os mesmos objetos já produzidos pelos
# módulos aprovados e exibidos no Streamlit.
#
# Abas:
#   1. Identificação
#   2. BP Ativo
#   3. BP Passivo
#   4. DRE
#   5. Indicadores
#   6. Relatório
#   7. Validação
#
# Saída:
#   bytes de um arquivo .xlsx, prontos para st.download_button.
# ============================================================


COR_PRIMARIA = "17365D"
COR_SECUNDARIA = "DCE6F1"
COR_CABECALHO = "244062"
COR_TEXTO_CLARO = "FFFFFF"
COR_BORDA = "B8C4CE"
COR_ALERTA = "FFF2CC"
COR_BLOQUEIO = "F4CCCC"
COR_OK = "D9EAD3"
COR_INFO = "D9EAF7"

INDICADORES_PERCENTUAIS = {
    "IPL",
    "PCT",
    "CE",
    "EFSAT",
    "RSV",
    "ROA",
    "ROE",
    *INDICADORES_FINANCEIROS,
}


def _valor_excel(valor: Any) -> Any:
    """Converte valores pandas/numpy em tipos seguros para openpyxl."""
    if valor is None:
        return None

    try:
        if pd.isna(valor):
            return "N/D"
    except (TypeError, ValueError):
        pass

    if hasattr(valor, "item"):
        try:
            valor = valor.item()
        except Exception:
            pass

    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()

    if isinstance(valor, (list, tuple, set, dict)):
        return str(valor)

    if isinstance(valor, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valor)

    return valor


def _nome_seguro(texto: str) -> str:
    texto = str(texto or "").strip()
    permitidos = []
    for caractere in texto:
        if caractere.isalnum():
            permitidos.append(caractere)
        elif caractere in {" ", "-", "_"}:
            permitidos.append("-")
    nome = "".join(permitidos)
    while "--" in nome:
        nome = nome.replace("--", "-")
    return nome.strip("-") or "empresa"


def nome_arquivo_exportacao(
    identificacao: dict[str, Any],
    anos: list[int] | tuple[int, ...] | None = None,
) -> str:
    empresa = (
        identificacao.get("DENOM_CIA")
        or identificacao.get("EMPRESA")
        or identificacao.get("NOME")
        or identificacao.get("CD_CVM")
        or "empresa"
    )

    empresa = _nome_seguro(str(empresa))[:55]

    if anos:
        anos_limpos = sorted({int(ano) for ano in anos})
        periodo = f"{anos_limpos[0]}-{anos_limpos[-1]}"
    else:
        periodo = "periodo"

    return f"Sistema-CVM-{empresa}-{periodo}.xlsx"


def _estilo_titulo(
    ws,
    titulo: str,
    subtitulo: str | None = None,
    coluna_final: int = 8,
) -> int:
    coluna_final = max(2, int(coluna_final))
    letra_final = get_column_letter(coluna_final)

    ws["A1"] = titulo
    ws["A1"].font = Font(
        name="Aptos Display",
        size=18,
        bold=True,
        color=COR_TEXTO_CLARO,
    )
    ws["A1"].fill = PatternFill("solid", fgColor=COR_PRIMARIA)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A1:{letra_final}1")

    linha = 2

    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(
            name="Aptos",
            size=10,
            italic=True,
            color="44546A",
        )
        ws["A2"].alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )
        ws.merge_cells(f"A2:{letra_final}2")
        ws.row_dimensions[2].height = 28
        linha = 4
    else:
        linha = 3

    return linha


def _formatar_tabela(
    ws,
    linha_cabecalho: int,
    linha_final: int,
    coluna_final: int,
) -> None:
    borda = Side(style="thin", color=COR_BORDA)

    for celula in ws[linha_cabecalho]:
        if celula.column > coluna_final:
            break
        celula.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=COR_TEXTO_CLARO,
        )
        celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        celula.border = Border(
            left=borda,
            right=borda,
            top=borda,
            bottom=borda,
        )

    for linha in ws.iter_rows(
        min_row=linha_cabecalho + 1,
        max_row=linha_final,
        min_col=1,
        max_col=coluna_final,
    ):
        for celula in linha:
            celula.font = Font(name="Aptos", size=10)
            celula.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            celula.border = Border(
                left=borda,
                right=borda,
                top=borda,
                bottom=borda,
            )

    ws.freeze_panes = f"A{linha_cabecalho + 1}"
    ws.auto_filter.ref = (
        f"A{linha_cabecalho}:"
        f"{get_column_letter(coluna_final)}{linha_final}"
    )


def _ajustar_larguras(ws, max_largura: int = 42) -> None:
    for coluna in range(1, ws.max_column + 1):
        maior = 0
        letra = get_column_letter(coluna)

        for linha in range(1, min(ws.max_row, 500) + 1):
            valor = ws.cell(linha, coluna).value
            if valor is None:
                continue
            texto = str(valor)
            if "\n" in texto:
                tamanho = max(len(parte) for parte in texto.splitlines())
            else:
                tamanho = len(texto)
            maior = max(maior, tamanho)

        largura = min(max(maior + 2, 11), max_largura)
        ws.column_dimensions[letra].width = largura


def _formato_numero_tabela(
    tabela: pd.DataFrame,
    indice_linha: int,
    nome_coluna: str,
) -> str:
    """Retorna apenas o formato visual; não altera o valor calculado."""
    nome = str(nome_coluna).upper().strip()

    if nome == "ORDEM" or nome == "ANO":
        return "0"

    if nome.startswith("VA_"):
        return '#,##0.00;(#,##0.00);0.00'

    # AV já está armazenada em pontos percentuais (ex.: 21,14).
    # O símbolo é literal para não multiplicar o valor por 100 outra vez.
    if nome.startswith("AV_"):
        return '0.00"%"'

    # AH já está armazenada como índice base fixa (ex.: 108,17).
    # O símbolo é literal para exibir 108,17% sem multiplicar por 100.
    if nome.startswith("AH_"):
        return '0.00"%"'

    indicador = ""
    unidade = ""

    if "INDICADOR" in tabela.columns:
        try:
            indicador = str(
                tabela.iloc[indice_linha]["INDICADOR"]
            ).strip().upper()
        except Exception:
            indicador = ""

    if "UNIDADE" in tabela.columns:
        try:
            unidade = str(
                tabela.iloc[indice_linha]["UNIDADE"]
            ).strip()
        except Exception:
            unidade = ""

    colunas_resultado = {
        "BASE",
        "INTERMEDIARIO",
        "RECENTE",
    }
    coluna_ano = nome.isdigit() and len(nome) == 4

    # Indicadores percentuais são armazenados como razão decimal
    # (ex.: 1,511 -> 151,10%). Aqui usamos o formato percentual do Excel.
    if (
        unidade == "%"
        or indicador in INDICADORES_PERCENTUAIS
    ) and (
        coluna_ano
        or nome in colunas_resultado
    ):
        return "0.00%"

    if coluna_ano or nome in colunas_resultado:
        return "0.00"

    return '#,##0.00;(#,##0.00);0.00'


def _escrever_dataframe(
    ws,
    df: pd.DataFrame | None,
    titulo: str,
    subtitulo: str | None = None,
) -> None:
    tabela = df.copy() if df is not None else None

    if tabela is not None:
        tabela.columns = [str(coluna) for coluna in tabela.columns]
        largura_tabela = max(8, len(tabela.columns))
    else:
        largura_tabela = 8

    inicio = _estilo_titulo(
        ws,
        titulo,
        subtitulo,
        coluna_final=largura_tabela,
    )

    if tabela is None or tabela.empty:
        ws.cell(inicio, 1, "Nenhum dado disponível.")
        ws.cell(inicio, 1).font = Font(
            name="Aptos",
            size=10,
            italic=True,
            color="666666",
        )
        return

    for j, coluna in enumerate(tabela.columns, start=1):
        ws.cell(inicio, j, _valor_excel(coluna))

    for indice_linha, linha in enumerate(
        tabela.itertuples(index=False, name=None)
    ):
        linha_excel = inicio + 1 + indice_linha

        for j, valor in enumerate(linha, start=1):
            celula = ws.cell(
                linha_excel,
                j,
                _valor_excel(valor),
            )

            if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                celula.number_format = _formato_numero_tabela(
                    tabela,
                    indice_linha,
                    tabela.columns[j - 1],
                )

    final = inicio + len(tabela)
    _formatar_tabela(ws, inicio, final, len(tabela.columns))
    _ajustar_larguras(ws)



def _escrever_identificacao(
    ws,
    identificacao: dict[str, Any],
    anos: list[int] | tuple[int, ...] | None,
    status_validacao: str | None,
) -> None:
    linha = _estilo_titulo(
        ws,
        "Sistema CVM — Identificação",
        "Exportação dos mesmos resultados utilizados no sistema de análise.",
    )

    campos = [
        (
            "Empresa",
            identificacao.get("DENOM_CIA")
            or identificacao.get("EMPRESA"),
        ),
        (
            "CD_CVM",
            identificacao.get("CD_CVM"),
        ),
        (
            "CNPJ",
            identificacao.get("CNPJ_CIA")
            or identificacao.get("CNPJ"),
        ),
    ]

    layout = identificacao.get(
        "LAYOUT_CVM"
    )

    if layout:
        campos.append(
            (
                "Layout CVM",
                layout,
            )
        )

    if identificacao.get(
        "TRATAMENTO_SETORIAL"
    ):
        campos.append(
            (
                "Tratamento setorial",
                (
                    "Sim — indicadores não comparáveis ao "
                    "modelo tradicional são exibidos como N/A."
                ),
            )
        )

    if anos:
        anos_limpos = sorted(
            {
                int(ano)
                for ano in anos
            }
        )

        campos.extend(
            [
                (
                    "Período analisado",
                    f"{anos_limpos[0]}–{anos_limpos[-1]}",
                ),
                (
                    "Ano-base da AH",
                    anos_limpos[0],
                ),
                (
                    "Exercício mais recente",
                    anos_limpos[-1],
                ),
            ]
        )

    campos.extend(
        [
            (
                "Status de validação",
                status_validacao or "N/D",
            ),
            (
                "Gerado em",
                datetime.now().strftime(
                    "%d/%m/%Y %H:%M:%S"
                ),
            ),
            (
                "Observação metodológica",
                (
                    "Este arquivo é uma exportação do Sistema CVM. "
                    "Não contém motor financeiro paralelo: demonstrações, "
                    "indicadores, relatório e validação provêm dos mesmos "
                    "resultados calculados pelo aplicativo."
                ),
            ),
        ]
    )

    borda = Side(
        style="thin",
        color=COR_BORDA,
    )

    for i, (campo, valor) in enumerate(
        campos,
        start=linha,
    ):
        ws.cell(
            i,
            1,
            campo,
        )
        ws.cell(
            i,
            2,
            _valor_excel(valor),
        )

        ws.cell(i, 1).font = Font(
            name="Aptos",
            bold=True,
            size=10,
        )

        ws.cell(i, 1).fill = PatternFill(
            "solid",
            fgColor=COR_SECUNDARIA,
        )

        for coluna in (1, 2):
            ws.cell(
                i,
                coluna,
            ).alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            ws.cell(
                i,
                coluna,
            ).border = Border(
                left=borda,
                right=borda,
                top=borda,
                bottom=borda,
            )

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 78


def _texto_relatorio(valor: Any) -> str:
    if valor is None:
        return "N/D"

    if isinstance(valor, pd.DataFrame):
        if valor.empty:
            return "N/D"
        return valor.to_string(index=False)

    if isinstance(valor, pd.Series):
        return valor.to_string()

    if isinstance(valor, dict):
        partes = []
        for chave, item in valor.items():
            if isinstance(item, (dict, list, tuple, pd.DataFrame, pd.Series)):
                continue
            partes.append(f"{chave}: {_texto_relatorio(item)}")
        return "\n".join(partes) if partes else str(valor)

    if isinstance(valor, (list, tuple)):
        if not valor:
            return "N/D"
        return "\n".join(
            f"• {_texto_relatorio(item)}"
            for item in valor
        )

    return ILLEGAL_CHARACTERS_RE.sub("", str(valor))


def _escrever_relatorio(
    ws,
    relatorio: Any,
) -> None:
    linha = _estilo_titulo(
        ws,
        "Relatório de Análise",
        "Interpretação automática exportada do mesmo relatório exibido no Streamlit.",
        coluna_final=12,
    )

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100

    if relatorio is None:
        ws.cell(linha, 1, "Relatório")
        ws.cell(linha, 2, "N/D")
        return

    borda = Side(style="thin", color=COR_BORDA)

    def escrever_secao(rotulo: str, conteudo: Any, nivel: int = 0) -> None:
        nonlocal linha

        if isinstance(conteudo, dict):
            ws.cell(linha, 1, rotulo)
            ws.cell(linha, 1).font = Font(
                name="Aptos",
                size=11 if nivel == 0 else 10,
                bold=True,
                color=COR_TEXTO_CLARO,
            )
            ws.cell(linha, 1).fill = PatternFill(
                "solid",
                fgColor=COR_CABECALHO if nivel == 0 else COR_PRIMARIA,
            )
            ws.merge_cells(
                start_row=linha,
                start_column=1,
                end_row=linha,
                end_column=2,
            )
            linha += 1

            for chave, valor in conteudo.items():
                escrever_secao(str(chave).replace("_", " ").title(), valor, nivel + 1)
            return

        if isinstance(conteudo, pd.DataFrame):
            ws.cell(linha, 1, rotulo)
            ws.cell(linha, 1).font = Font(
                name="Aptos",
                bold=True,
                size=10,
            )
            ws.cell(linha, 1).fill = PatternFill("solid", fgColor=COR_SECUNDARIA)
            linha += 1

            if conteudo.empty:
                ws.cell(linha, 1, "N/D")
                linha += 2
                return

            colunas = [str(c) for c in conteudo.columns]
            for j, coluna in enumerate(colunas, start=1):
                ws.cell(linha, j, coluna)

            cab = linha
            linha += 1

            tabela_relatorio = conteudo.copy()
            tabela_relatorio.columns = [
                str(coluna) for coluna in tabela_relatorio.columns
            ]

            for indice_linha, dados in enumerate(
                tabela_relatorio.itertuples(index=False, name=None)
            ):
                for j, valor in enumerate(dados, start=1):
                    celula = ws.cell(linha, j, _valor_excel(valor))

                    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                        celula.number_format = _formato_numero_tabela(
                            tabela_relatorio,
                            indice_linha,
                            tabela_relatorio.columns[j - 1],
                        )
                linha += 1

            _formatar_tabela(
                ws,
                cab,
                linha - 1,
                len(colunas),
            )
            linha += 1
            return

        ws.cell(linha, 1, rotulo)
        ws.cell(linha, 2, _texto_relatorio(conteudo))
        ws.cell(linha, 1).font = Font(name="Aptos", bold=True, size=10)
        ws.cell(linha, 1).fill = PatternFill("solid", fgColor=COR_SECUNDARIA)

        for coluna in (1, 2):
            ws.cell(linha, coluna).alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            ws.cell(linha, coluna).border = Border(
                left=borda,
                right=borda,
                top=borda,
                bottom=borda,
            )

        texto = str(ws.cell(linha, 2).value or "")
        ws.row_dimensions[linha].height = min(
            max(18, 15 * (texto.count("\n") + 1)),
            120,
        )
        linha += 1

    if isinstance(relatorio, dict):
        for chave, conteudo in relatorio.items():
            escrever_secao(
                str(chave).replace("_", " ").title(),
                conteudo,
                0,
            )
    else:
        escrever_secao("Relatório", relatorio, 0)

    _ajustar_larguras(ws, max_largura=55)


def _colorir_status_validacao(ws) -> None:
    cabecalhos = {
        str(ws.cell(4, coluna).value).strip().upper(): coluna
        for coluna in range(1, ws.max_column + 1)
        if ws.cell(4, coluna).value is not None
    }

    coluna_status = cabecalhos.get("STATUS")

    if not coluna_status:
        for linha in range(1, min(ws.max_row, 15) + 1):
            for coluna in range(1, ws.max_column + 1):
                if str(ws.cell(linha, coluna).value).strip().upper() == "STATUS":
                    coluna_status = coluna
                    inicio = linha + 1
                    break
            if coluna_status:
                break
        else:
            return
    else:
        inicio = 5

    fills = {
        "OK": COR_OK,
        "INFO": COR_INFO,
        "ALERTA": COR_ALERTA,
        "BLOQUEIO": COR_BLOQUEIO,
    }

    for linha in range(inicio, ws.max_row + 1):
        celula = ws.cell(linha, coluna_status)
        status = str(celula.value or "").strip().upper()
        cor = fills.get(status)
        if cor:
            celula.fill = PatternFill("solid", fgColor=cor)
            celula.font = Font(name="Aptos", bold=True, size=10)


def _validacao_financeira_exportacao(
    validacao: pd.DataFrame | None,
    relatorio: Any,
) -> pd.DataFrame | None:
    if validacao is None:
        return None

    resultado = validacao.copy()

    if {
        "CATEGORIA",
        "TESTE",
    }.issubset(resultado.columns):
        categoria = resultado["CATEGORIA"].astype(str)
        teste = resultado["TESTE"].astype(str)
        contrato_antigo = (
            categoria.eq("Indicadores")
            & teste.isin(
                [
                    "Quantidade de indicadores",
                    "Cobertura dos indicadores",
                ]
            )
        )

        resultado = resultado.loc[
            ~(
                contrato_antigo
                | categoria.eq("DuPont")
            )
        ].copy()

    base_financeira = (
        relatorio.get("INDICADORES_SETORIAIS")
        if isinstance(relatorio, dict)
        else None
    )

    if (
        not isinstance(base_financeira, pd.DataFrame)
        or base_financeira.empty
        or not {"ANO", "INDICADOR", "STATUS"}.issubset(
            base_financeira.columns
        )
    ):
        return resultado

    linhas = []
    esperados = len(INDICADORES_FINANCEIROS) + 2

    for ano in sorted(
        base_financeira["ANO"].dropna().astype(int).unique()
    ):
        dados_ano = base_financeira.loc[
            base_financeira["ANO"].eq(ano)
        ]
        qtd = int(dados_ano["INDICADOR"].nunique())

        linhas.append(
            {
                "CATEGORIA": "Indicadores",
                "ANO": ano,
                "TESTE": "Quantidade de indicadores",
                "STATUS": (
                    "OK"
                    if qtd == esperados
                    else "BLOQUEIO"
                ),
                "DETALHE": (
                    f"{qtd} indicadores financeiros "
                    "calculados/representados."
                    if qtd == esperados
                    else (
                        f"{qtd} indicadores financeiros; "
                        f"esperado {esperados}."
                    )
                ),
            }
        )

        status = dados_ano["STATUS"].astype(str).str.strip()

        nd = (
            dados_ano.loc[status.eq("N/D"), "INDICADOR"]
            .astype(str).tolist()
        )
        na = (
            dados_ano.loc[status.eq("N/A"), "INDICADOR"]
            .astype(str).tolist()
        )

        detalhes = []

        if nd:
            detalhes.append(
                "N/D entre indicadores financeiros: "
                + ", ".join(nd)
            )

        if na:
            detalhes.append(
                "N/A entre indicadores financeiros: "
                + ", ".join(na)
            )

        linhas.append(
            {
                "CATEGORIA": "Indicadores",
                "ANO": ano,
                "TESTE": "Cobertura dos indicadores",
                "STATUS": (
                    "ALERTA"
                    if nd
                    else (
                        "INFO"
                        if na
                        else "OK"
                    )
                ),
                "DETALHE": (
                    "; ".join(detalhes)
                    if detalhes
                    else (
                        f"Todos os {esperados} indicadores financeiros "
                        "possuem resultado numérico."
                    )
                ),
            }
        )

    cobertura = pd.DataFrame(linhas).reindex(
        columns=resultado.columns
    )

    return pd.concat(
        [resultado, cobertura],
        ignore_index=True,
    )


def gerar_excel_sistema(
    *,
    identificacao: dict[str, Any],
    bp_ativo: pd.DataFrame | None,
    bp_passivo: pd.DataFrame | None,
    dre: pd.DataFrame | None,
    indicadores: pd.DataFrame | None,
    relatorio: Any,
    validacao: pd.DataFrame | None,
    anos: list[int] | tuple[int, ...] | None = None,
    status_validacao: str | None = None,
) -> bytes:
    """
    Gera o arquivo Excel completo do Sistema CVM.

    Importante:
    - não recalcula AV, AH, indicadores ou relatório;
    - apenas exporta os mesmos objetos já produzidos pelo app;
    - retorna bytes prontos para uso no Streamlit.
    """
    layout_financeiro = (
        LAYOUT_FINANCEIRA
        in str(
            identificacao.get(
                "LAYOUT_CVM",
                "",
            )
        )
    )

    wb = Workbook()

    # Remove planilha padrão.
    ws_padrao = wb.active
    wb.remove(ws_padrao)

    ws_id = wb.create_sheet("Identificação")
    _escrever_identificacao(
        ws_id,
        identificacao=identificacao,
        anos=anos,
        status_validacao=status_validacao,
    )

    ws_ativo = wb.create_sheet("BP Ativo")
    _escrever_dataframe(
        ws_ativo,
        bp_ativo,
        "Balanço Patrimonial — Ativo",
        "Valores e análises exportados do mesmo quadro utilizado pelo sistema.",
    )

    ws_passivo = wb.create_sheet("BP Passivo")
    _escrever_dataframe(
        ws_passivo,
        bp_passivo,
        "Balanço Patrimonial — Passivo e Patrimônio Líquido",
        "Valores e análises exportados do mesmo quadro utilizado pelo sistema.",
    )

    dre_exportacao = (
        dre.copy()
        if dre is not None
        else None
    )

    if (
        dre_exportacao is not None
        and bool(
            getattr(
                dre,
                "attrs",
                {},
            ).get(
                "AV_NAO_APLICAVEL",
                False,
            )
        )
    ):
        for coluna in dre_exportacao.columns:
            if str(coluna).startswith("AV_"):
                dre_exportacao[coluna] = "N/A"

    ws_dre = wb.create_sheet("DRE")
    _escrever_dataframe(
        ws_dre,
        dre_exportacao,
        "Demonstração do Resultado do Exercício",
        "Valores e análises exportados do mesmo quadro utilizado pelo sistema.",
    )

    if layout_financeiro:
        analise_financeira = (
            relatorio.get(
                "ANALISE_INDICADORES"
            )
            if isinstance(
                relatorio,
                dict,
            )
            else None
        )

        indicadores_exportacao = (
            analise_financeira.copy()
            if isinstance(
                analise_financeira,
                pd.DataFrame,
            )
            else None
        )
    else:
        indicadores_exportacao = (
            indicadores.copy()
            if indicadores is not None
            else None
        )

    if (
        indicadores_exportacao is not None
        and not indicadores_exportacao.empty
        and "APLICABILIDADE"
        in indicadores_exportacao.columns
    ):
        mascara_na = (
            indicadores_exportacao[
                "APLICABILIDADE"
            ]
            .astype(str)
            .eq("NAO_APLICAVEL")
        )

        for coluna in indicadores_exportacao.columns:
            if (
                isinstance(coluna, int)
                or coluna in {
                    "BASE",
                    "INTERMEDIARIO",
                    "RECENTE",
                }
            ):
                # Pandas 3 não permite gravar texto diretamente
                # em coluna float64. Convertemos somente as colunas
                # de valores exportados para object antes do N/A.
                indicadores_exportacao[
                    coluna
                ] = indicadores_exportacao[
                    coluna
                ].astype("object")

                if (
                    layout_financeiro
                    and coluna in {
                        "BASE",
                        "INTERMEDIARIO",
                        "RECENTE",
                    }
                ):
                    indicadores_exportacao.loc[
                        indicadores_exportacao[
                            coluna
                        ].isna(),
                        coluna,
                    ] = "N/D"

                indicadores_exportacao.loc[
                    mascara_na,
                    coluna,
                ] = "N/A"

    ws_ind = wb.create_sheet("Indicadores")
    _escrever_dataframe(
        ws_ind,
        indicadores_exportacao,
        "Indicadores Financeiros",
        (
            "Os 9 indicadores aplicáveis ao layout FINANCEIRA "
            "seguem os resultados já calculados pelos motores aprovados."
            if layout_financeiro
            else (
                "Os 12 indicadores seguem os resultados já calculados "
                "pelo motor financeiro aprovado."
            )
        ),
    )

    validacao_exportacao = (
        _validacao_financeira_exportacao(
            validacao,
            relatorio,
        )
        if layout_financeiro
        else validacao
    )

    relatorio_exportacao = relatorio

    if layout_financeiro and isinstance(
        relatorio,
        dict,
    ):
        relatorio_exportacao = {
            chave: valor
            for chave, valor in relatorio.items()
            if chave != "DUPONT"
        }

        if "VALIDACAO" in relatorio_exportacao:
            relatorio_exportacao[
                "VALIDACAO"
            ] = validacao_exportacao

        if (
            "ANALISE_INDICADORES" in relatorio_exportacao
            and isinstance(
                indicadores_exportacao,
                pd.DataFrame,
            )
        ):
            relatorio_exportacao[
                "ANALISE_INDICADORES"
            ] = indicadores_exportacao.copy()

    ws_rel = wb.create_sheet("Relatório")
    _escrever_relatorio(
        ws_rel,
        relatorio_exportacao,
    )

    ws_val = wb.create_sheet("Validação")
    _escrever_dataframe(
        ws_val,
        validacao_exportacao,
        "Validação / Auditoria",
        "BLOQUEIO suspende interpretação conclusiva; ALERTA permite análise com ressalva.",
    )
    _colorir_status_validacao(ws_val)

    # Configurações gerais.
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5

    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    return arquivo.getvalue()


__all__ = [
    "gerar_excel_sistema",
    "nome_arquivo_exportacao",
]
