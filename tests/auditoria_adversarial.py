from __future__ import annotations

import argparse
import sys
import time
import traceback
import zipfile
from io import BytesIO
from openpyxl import load_workbook
from pathlib import Path
from xml.etree import ElementTree as ET

import numpy as np
import pandas as pd


# ============================================================
# AUDITORIA ADVERSARIAL — SISTEMA CVM
# ============================================================
#
# Uso:
#   python tests\auditoria_adversarial.py --modo rapido
#   python tests\auditoria_adversarial.py --modo completo
#
# O modo rápido:
#   - audita o catálogo inteiro;
#   - identifica automaticamente casos extremos;
#   - executa testes profundos em uma amostra adversarial.
#
# O modo completo:
#   - faz tudo do modo rápido;
#   - executa o núcleo financeiro em TODAS as companhias;
#   - mantém testes pesados (relatório/dashboard/exportação)
#     na amostra adversarial para não multiplicar o tempo.
#
# Não altera cálculos nem dados do sistema.
# Apenas grava CSVs de auditoria em data/processed.
# ============================================================


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

DATA_DIR = ROOT / "data" / "processed"
ARQUIVO_EMPRESAS = DATA_DIR / "empresas.parquet"
ARQUIVO_DFP = DATA_DIR / "dfp_2023_2025.parquet"

INDICADORES_ESPERADOS = [
    "IPL",
    "PCT",
    "CE",
    "EFSAT",
    "LG",
    "LC",
    "LS",
    "ICJ",
    "GA",
    "RSV",
    "ROA",
    "ROE",
]

ABAS_EXPORTACAO = [
    "Identificação",
    "BP Ativo",
    "BP Passivo",
    "DRE",
    "Indicadores",
    "Relatório",
    "Validação",
]

STATUS_VALIDOS = {
    "OK",
    "INFO",
    "ALERTA",
    "BLOQUEIO",
}

CHAVES_RELATORIO = {
    "STATUS_VALIDACAO",
    "RESUMO_EXECUTIVO",
    "PRINCIPAIS_MUDANCAS",
    "ANALISE_INDICADORES",
    "SINTESES_GRUPOS",
    "PONTOS_ATENCAO",
    "CONCLUSAO",
}

CHAVES_GRAFICOS = {
    "estrutura",
    "liquidez",
    "icj",
    "rentabilidade",
    "ga",
}


# ------------------------------------------------------------
# IMPORTS DO PROJETO
# ------------------------------------------------------------

try:
    from src.demonstracoes import resumo_empresa
    from src.padronizacao import montar_demonstracao_apresentacao
    from src.layouts_cvm import (
        LAYOUT_PADRAO,
        LAYOUT_FINANCEIRA,
        LAYOUT_SEGUROS_ESPECIAL,
    )
    from src.indicadores import (
        quadro_indicadores,
        validar_dupont,
    )
    from src.validacao import (
        validar_empresa,
        status_geral,
    )
    from src.relatorio import gerar_relatorio
    from src.graficos import (
        montar_cards_dashboard,
        montar_graficos_dashboard,
    )
    from src.exportacao import (
        gerar_excel_sistema,
        nome_arquivo_exportacao,
    )
except Exception as exc:
    print("=" * 78)
    print("FALHA AO IMPORTAR O SISTEMA")
    print("=" * 78)
    print(repr(exc))
    print()
    print("Execute este arquivo a partir da pasta raiz do projeto:")
    print(r"python tests\auditoria_adversarial.py --modo rapido")
    raise


# ------------------------------------------------------------
# UTILIDADES
# ------------------------------------------------------------

def normalizar_cd(valor) -> str:
    return str(valor).strip().zfill(6)


def eh_numero_finito(valor) -> bool:
    if valor is None or pd.isna(valor):
        return True
    try:
        return bool(np.isfinite(float(valor)))
    except (TypeError, ValueError):
        return True


def adicionar_resultado(
    resultados: list[dict],
    cd_cvm: str,
    empresa: str,
    teste: str,
    status: str,
    detalhe: str = "",
    caso: str = "",
) -> None:
    resultados.append(
        {
            "CD_CVM": cd_cvm,
            "EMPRESA": empresa,
            "CASO": caso,
            "TESTE": teste,
            "STATUS": status,
            "DETALHE": detalhe,
        }
    )


def nomes_abas_xlsx(conteudo: bytes) -> list[str]:
    """Lê nomes de abas diretamente do XML do .xlsx."""
    with zipfile.ZipFile(BytesIO(conteudo), "r") as z:
        xml = z.read("xl/workbook.xml")

    raiz = ET.fromstring(xml)
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

    return [
        elem.attrib["name"]
        for elem in raiz.findall(".//m:sheets/m:sheet", ns)
    ]


def carregar_catalogo() -> pd.DataFrame:
    if not ARQUIVO_EMPRESAS.exists():
        raise FileNotFoundError(
            f"Catálogo não encontrado: {ARQUIVO_EMPRESAS}"
        )

    df = pd.read_parquet(ARQUIVO_EMPRESAS).copy()

    obrigatorias = {
        "CD_CVM",
        "CNPJ_CIA",
        "DENOM_CIA",
    }

    faltantes = obrigatorias.difference(df.columns)
    if faltantes:
        raise ValueError(
            "Catálogo sem colunas obrigatórias: "
            + ", ".join(sorted(faltantes))
        )

    df["CD_CVM"] = (
        df["CD_CVM"]
        .astype("string")
        .str.strip()
        .str.zfill(6)
    )

    df["CNPJ_CIA"] = (
        df["CNPJ_CIA"]
        .astype("string")
        .str.replace(r"\D", "", regex=True)
        .str.zfill(14)
    )

    df["DENOM_CIA"] = (
        df["DENOM_CIA"]
        .astype("string")
        .str.strip()
    )

    return df.reset_index(drop=True)


def auditar_catalogo(
    catalogo: pd.DataFrame,
    resultados: list[dict],
) -> pd.DataFrame:
    duplicados_cd = catalogo[
        catalogo["CD_CVM"].duplicated(keep=False)
    ]

    duplicados_cnpj = catalogo[
        catalogo["CNPJ_CIA"].duplicated(keep=False)
    ]

    homonimos = catalogo[
        catalogo["DENOM_CIA"].duplicated(keep=False)
    ].sort_values(
        ["DENOM_CIA", "CD_CVM"]
    )

    adicionar_resultado(
        resultados,
        "",
        "CATÁLOGO",
        "CD_CVM único",
        "OK" if duplicados_cd.empty else "FALHA",
        (
            "0 duplicidades"
            if duplicados_cd.empty
            else f"{len(duplicados_cd)} linhas com CD_CVM duplicado"
        ),
        "estrutura",
    )

    adicionar_resultado(
        resultados,
        "",
        "CATÁLOGO",
        "CNPJ único",
        "OK" if duplicados_cnpj.empty else "FALHA",
        (
            "0 duplicidades"
            if duplicados_cnpj.empty
            else f"{len(duplicados_cnpj)} linhas com CNPJ duplicado"
        ),
        "estrutura",
    )

    adicionar_resultado(
        resultados,
        "",
        "CATÁLOGO",
        "Homônimos identificáveis",
        "OK",
        (
            f"{homonimos['DENOM_CIA'].nunique()} denominação(ões) "
            "repetida(s); o sistema usa CD_CVM/CNPJ para distinguir."
        ),
        "homônimos",
    )

    return homonimos


# ------------------------------------------------------------
# DESCOBERTA AUTOMÁTICA DE CASOS ADVERSARIAIS
# ------------------------------------------------------------

def _preparar_dados_extremos() -> pd.DataFrame:
    if not ARQUIVO_DFP.exists():
        return pd.DataFrame()

    colunas = [
        "CD_CVM",
        "ANO",
        "DEMONSTRACAO",
        "CD_CONTA",
        "VL_CONTA",
    ]

    df = pd.read_parquet(
        ARQUIVO_DFP,
        columns=colunas,
    ).copy()

    df["CD_CVM"] = (
        df["CD_CVM"]
        .astype("string")
        .str.strip()
        .str.zfill(6)
    )

    df["CD_CONTA"] = (
        df["CD_CONTA"]
        .astype("string")
        .str.strip()
    )

    df["DEMONSTRACAO"] = (
        df["DEMONSTRACAO"]
        .astype("string")
        .str.strip()
        .str.upper()
    )

    df["VL_CONTA"] = pd.to_numeric(
        df["VL_CONTA"],
        errors="coerce",
    )

    df["ANO"] = pd.to_numeric(
        df["ANO"],
        errors="coerce",
    ).astype("Int64")

    return df


def descobrir_casos(
    catalogo: pd.DataFrame,
    homonimos: pd.DataFrame,
) -> pd.DataFrame:
    casos: list[dict] = []

    def incluir(cd: str, caso: str, detalhe: str = "") -> None:
        cd = normalizar_cd(cd)
        linha = catalogo.loc[
            catalogo["CD_CVM"].eq(cd)
        ]

        if linha.empty:
            return

        casos.append(
            {
                "CD_CVM": cd,
                "EMPRESA": str(linha.iloc[0]["DENOM_CIA"]),
                "CASO": caso,
                "DETALHE_CASO": detalhe,
            }
        )

    # Benchmark obrigatório já validado.
    if catalogo["CD_CVM"].eq("004170").any():
        incluir(
            "004170",
            "benchmark",
            "VALE S.A. — referência de regressão",
        )

    # Regressões setoriais obrigatórias.
    # Estes IDs são usados somente na auditoria/teste de regressão,
    # nunca na lógica de classificação do aplicativo.
    amostras_setoriais = [
        (
            "019348",
            "layout financeiro",
            "Itaú — referência FINANCEIRA",
        ),
        (
            "000906",
            "layout financeiro",
            "Bradesco — referência FINANCEIRA",
        ),
        (
            "023159",
            "layout seguros especial",
            "BB Seguridade — referência SEGUROS_ESPECIAL",
        ),
        (
            "024180",
            "layout seguros especial",
            "IRB — referência SEGUROS_ESPECIAL",
        ),
    ]

    for cd, caso_setorial, detalhe_setorial in amostras_setoriais:
        if catalogo["CD_CVM"].eq(cd).any():
            incluir(
                cd,
                caso_setorial,
                detalhe_setorial,
            )

    # Homônimos reais do catálogo.
    for _, linha in homonimos.head(6).iterrows():
        incluir(
            linha["CD_CVM"],
            "homônimo",
            "Denominação repetida no catálogo",
        )

    dados = _preparar_dados_extremos()

    if not dados.empty:
        chaves = {
            "PL": ("BPP", "2.03"),
            "LL": ("DRE", "3.11"),
            "AT": ("BPA", "1"),
            "PC": ("BPP", "2.01"),
            "PNC": ("BPP", "2.02"),
            "RECEITA": ("DRE", "3.01"),
            "DESP_FIN": ("DRE", "3.06.02"),
        }

        tabelas = {}

        for nome, (dem, conta) in chaves.items():
            parte = dados.loc[
                dados["DEMONSTRACAO"].eq(dem)
                & dados["CD_CONTA"].eq(conta),
                [
                    "CD_CVM",
                    "ANO",
                    "VL_CONTA",
                ],
            ].copy()

            parte = (
                parte
                .sort_values(["CD_CVM", "ANO"])
                .drop_duplicates(
                    ["CD_CVM", "ANO"],
                    keep="last",
                )
            )

            parte = parte.rename(
                columns={"VL_CONTA": nome}
            )

            tabelas[nome] = parte

        # PL negativo.
        pl = tabelas["PL"]
        neg_pl = pl.loc[
            pl["PL"].lt(0)
        ].head(3)

        for _, linha in neg_pl.iterrows():
            incluir(
                linha["CD_CVM"],
                "PL negativo",
                f"{int(linha['ANO'])}: PL={linha['PL']}",
            )

        # Prejuízo.
        ll = tabelas["LL"]
        prejuizo = ll.loc[
            ll["LL"].lt(0)
        ].head(3)

        for _, linha in prejuizo.iterrows():
            incluir(
                linha["CD_CVM"],
                "prejuízo",
                f"{int(linha['ANO'])}: LL={linha['LL']}",
            )

        # Denominadores zero encontrados diretamente.
        for nome in [
            "PL",
            "AT",
            "PC",
            "RECEITA",
            "DESP_FIN",
        ]:
            parte = tabelas[nome]
            zeros = parte.loc[
                parte[nome].eq(0)
            ].head(2)

            for _, linha in zeros.iterrows():
                incluir(
                    linha["CD_CVM"],
                    f"{nome} zero",
                    f"{int(linha['ANO'])}: {nome}=0",
                )

        # Capital de terceiros = PC + PNC igual a zero.
        ct = tabelas["PC"].merge(
            tabelas["PNC"],
            on=["CD_CVM", "ANO"],
            how="outer",
        )

        ct["CT"] = (
            pd.to_numeric(ct["PC"], errors="coerce")
            + pd.to_numeric(ct["PNC"], errors="coerce")
        )

        for _, linha in ct.loc[
            ct["CT"].eq(0)
        ].head(2).iterrows():
            incluir(
                linha["CD_CVM"],
                "capital de terceiros zero",
                f"{int(linha['ANO'])}: PC+PNC=0",
            )

        # Mudança de sinal do lucro líquido.
        piv_ll = ll.pivot_table(
            index="CD_CVM",
            columns="ANO",
            values="LL",
            aggfunc="last",
        )

        for cd, linha in piv_ll.iterrows():
            valores = linha.dropna().sort_index()
            sinais = np.sign(valores.to_numpy(dtype=float))

            if len(sinais) >= 2 and np.any(
                sinais[1:] * sinais[:-1] < 0
            ):
                incluir(
                    cd,
                    "mudança de sinal do resultado",
                    "Lucro/prejuízo muda de sinal entre exercícios",
                )

                if sum(
                    item["CASO"] == "mudança de sinal do resultado"
                    for item in casos
                ) >= 3:
                    break

        # Mudança de sinal do PL.
        piv_pl = pl.pivot_table(
            index="CD_CVM",
            columns="ANO",
            values="PL",
            aggfunc="last",
        )

        for cd, linha in piv_pl.iterrows():
            valores = linha.dropna().sort_index()
            sinais = np.sign(valores.to_numpy(dtype=float))

            if len(sinais) >= 2 and np.any(
                sinais[1:] * sinais[:-1] < 0
            ):
                incluir(
                    cd,
                    "mudança de sinal do PL",
                    "PL muda de sinal entre exercícios",
                )

                if sum(
                    item["CASO"] == "mudança de sinal do PL"
                    for item in casos
                ) >= 3:
                    break

        # Empresas com menos de três anos efetivamente presentes.
        cobertura = (
            dados.groupby("CD_CVM")["ANO"]
            .nunique()
            .sort_values()
        )

        for cd, n in cobertura.loc[
            cobertura.lt(3)
        ].head(3).items():
            incluir(
                cd,
                "período incompleto",
                f"Apenas {int(n)} exercício(s) em 2023–2025",
            )

    # Completa a amostra com empresas distribuídas pelo catálogo.
    if len(catalogo) > 0:
        indices = sorted(
            set(
                np.linspace(
                    0,
                    len(catalogo) - 1,
                    num=min(8, len(catalogo)),
                    dtype=int,
                )
            )
        )

        for indice in indices:
            linha = catalogo.iloc[int(indice)]
            incluir(
                linha["CD_CVM"],
                "amostra distribuída",
                "Amostra sistemática do catálogo",
            )

    df = pd.DataFrame(casos)

    if df.empty:
        return df

    # Um mesmo CD pode representar vários casos; mantém todos os rótulos
    # em uma única linha para a execução profunda.
    consolidado = (
        df.groupby(
            ["CD_CVM", "EMPRESA"],
            as_index=False,
        )
        .agg(
            CASO=(
                "CASO",
                lambda s: " | ".join(
                    dict.fromkeys(str(x) for x in s)
                ),
            ),
            DETALHE_CASO=(
                "DETALHE_CASO",
                lambda s: " | ".join(
                    dict.fromkeys(
                        str(x)
                        for x in s
                        if str(x).strip()
                    )
                ),
            ),
        )
        .sort_values(
            ["CASO", "EMPRESA", "CD_CVM"]
        )
        .reset_index(drop=True)
    )

    # Limita a amostra profunda para manter o teste rápido utilizável.
    # Casos extremos ficam antes da amostra genérica.
    prioridade = consolidado["CASO"].str.contains(
        "benchmark|layout|homônimo|negativo|prejuízo|zero|mudança|incompleto",
        case=False,
        regex=True,
    )

    prioritarios = consolidado.loc[prioridade]
    genericos = consolidado.loc[~prioridade]

    return pd.concat(
        [
            prioritarios.head(24),
            genericos.head(8),
        ],
        ignore_index=True,
    )


# ------------------------------------------------------------
# TESTE DE UMA COMPANHIA
# ------------------------------------------------------------

def testar_empresa(
    cd_cvm: str,
    empresa: str,
    caso: str,
    resultados: list[dict],
    profundo: bool,
) -> None:
    inicio = time.perf_counter()

    try:
        resumo = resumo_empresa(cd_cvm)
        disponiveis = sorted(
            int(ano)
            for ano in resumo.get(
                "ANOS_DISPONIVEIS",
                [],
            )
        )

        if not disponiveis:
            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                "Exercícios disponíveis",
                "FALHA",
                "Nenhum exercício encontrado.",
                caso,
            )
            return

        recente = max(disponiveis)
        janela = [
            ano
            for ano in [
                recente - 2,
                recente - 1,
                recente,
            ]
            if ano in disponiveis
        ]

        if not janela:
            janela = disponiveis[-3:]

        adicionar_resultado(
            resultados,
            cd_cvm,
            empresa,
            "Seleção de período",
            "OK",
            f"Anos usados: {janela}",
            caso,
        )

        demonstracoes = {}

        for dem in [
            "BPA",
            "BPP",
            "DRE",
        ]:
            quadro = montar_demonstracao_apresentacao(
                cd_cvm=cd_cvm,
                demonstracao=dem,
                anos=janela,
            )

            demonstracoes[dem] = quadro

            if quadro.empty:
                adicionar_resultado(
                    resultados,
                    cd_cvm,
                    empresa,
                    f"{dem} padronizada",
                    "FALHA",
                    "Quadro vazio.",
                    caso,
                )
                continue

            colunas_necessarias = {
                "LINHA",
                "TIPO",
            }

            for ano in janela:
                colunas_necessarias.update(
                    {
                        f"VA_{ano}",
                        f"AV_{ano}",
                        f"AH_{ano}",
                    }
                )

            ausentes = (
                colunas_necessarias
                .difference(quadro.columns)
            )

            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                f"{dem} padronizada",
                "OK" if not ausentes else "FALHA",
                (
                    f"{len(quadro)} linhas"
                    if not ausentes
                    else "Colunas ausentes: "
                    + ", ".join(sorted(ausentes))
                ),
                caso,
            )

            numericas = [
                coluna
                for coluna in quadro.columns
                if str(coluna).startswith(
                    ("VA_", "AV_", "AH_")
                )
            ]

            infinitos = 0

            for coluna in numericas:
                serie = pd.to_numeric(
                    quadro[coluna],
                    errors="coerce",
                )

                infinitos += int(
                    np.isinf(
                        serie.to_numpy(
                            dtype=float,
                            na_value=np.nan,
                        )
                    ).sum()
                )

            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                f"{dem} sem infinito",
                "OK" if infinitos == 0 else "FALHA",
                f"{infinitos} valor(es) infinito(s)",
                caso,
            )

        indicadores = quadro_indicadores(
            cd_cvm,
            anos=janela,
        )

        codigos = (
            indicadores["INDICADOR"]
            .astype(str)
            .tolist()
            if not indicadores.empty
            and "INDICADOR" in indicadores.columns
            else []
        )

        faltam = [
            codigo
            for codigo in INDICADORES_ESPERADOS
            if codigo not in codigos
        ]

        duplicados = (
            indicadores["INDICADOR"]
            .duplicated()
            .any()
            if not indicadores.empty
            and "INDICADOR" in indicadores.columns
            else False
        )

        adicionar_resultado(
            resultados,
            cd_cvm,
            empresa,
            "12 indicadores representados",
            (
                "OK"
                if not faltam
                and not duplicados
                and len(codigos) == 12
                else "FALHA"
            ),
            (
                f"{len(codigos)} linhas; "
                f"faltantes={faltam or 'nenhum'}; "
                f"duplicados={bool(duplicados)}"
            ),
            caso,
        )

        # ----------------------------------------------------
        # LAYOUT E APLICABILIDADE
        # ----------------------------------------------------

        layouts = (
            indicadores["LAYOUT"]
            .dropna()
            .astype(str)
            .unique()
            .tolist()
            if (
                not indicadores.empty
                and "LAYOUT" in indicadores.columns
            )
            else []
        )

        layout_codigo = (
            layouts[0]
            if len(layouts) == 1
            else None
        )

        aplicabilidades = (
            indicadores["APLICABILIDADE"]
            .astype(str)
            if (
                not indicadores.empty
                and "APLICABILIDADE" in indicadores.columns
            )
            else pd.Series(dtype="string")
        )

        qtd_aplicaveis = int(
            aplicabilidades
            .eq("APLICAVEL")
            .sum()
        )

        qtd_na = int(
            aplicabilidades
            .eq("NAO_APLICAVEL")
            .sum()
        )

        esperados_por_layout = {
            LAYOUT_PADRAO: (12, 0),
            LAYOUT_FINANCEIRA: (2, 10),
            LAYOUT_SEGUROS_ESPECIAL: (2, 10),
            "OUTRO": (0, 12),
        }

        esperado = (
            esperados_por_layout.get(
                layout_codigo
            )
        )

        layout_ok = (
            esperado is not None
            and (
                qtd_aplicaveis,
                qtd_na,
            )
            == esperado
        )

        adicionar_resultado(
            resultados,
            cd_cvm,
            empresa,
            "Layout e aplicabilidade",
            "OK" if layout_ok else "FALHA",
            (
                f"layout={layout_codigo}; "
                f"aplicáveis={qtd_aplicaveis}; "
                f"N/A={qtd_na}; "
                f"esperado={esperado}"
            ),
            caso,
        )

        tratamento_setorial = (
            layout_codigo
            in {
                LAYOUT_FINANCEIRA,
                LAYOUT_SEGUROS_ESPECIAL,
            }
        )

        if tratamento_setorial:
            dre_apresentacao = (
                demonstracoes.get(
                    "DRE"
                )
            )

            colunas_av = (
                [
                    coluna
                    for coluna
                    in dre_apresentacao.columns
                    if str(
                        coluna
                    ).startswith(
                        "AV_"
                    )
                ]
                if isinstance(
                    dre_apresentacao,
                    pd.DataFrame,
                )
                else []
            )

            av_na = bool(
                getattr(
                    dre_apresentacao,
                    "attrs",
                    {},
                ).get(
                    "AV_NAO_APLICAVEL",
                    False,
                )
            )

            av_sem_valores = bool(
                colunas_av
                and all(
                    dre_apresentacao[
                        coluna
                    ].isna().all()
                    for coluna
                    in colunas_av
                )
            )

            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                "DRE setorial — AV não aplicável",
                (
                    "OK"
                    if (
                        av_na
                        and av_sem_valores
                    )
                    else "FALHA"
                ),
                (
                    f"flag={av_na}; "
                    f"colunas AV={len(colunas_av)}; "
                    f"todas vazias={av_sem_valores}"
                ),
                caso,
            )

        infinitos_ind = 0

        for ano in janela:
            if ano in indicadores.columns:
                serie = pd.to_numeric(
                    indicadores[ano],
                    errors="coerce",
                )
                infinitos_ind += int(
                    np.isinf(
                        serie.to_numpy(
                            dtype=float,
                            na_value=np.nan,
                        )
                    ).sum()
                )

        adicionar_resultado(
            resultados,
            cd_cvm,
            empresa,
            "Indicadores sem infinito",
            "OK" if infinitos_ind == 0 else "FALHA",
            f"{infinitos_ind} valor(es) infinito(s)",
            caso,
        )

        dupont = validar_dupont(
            cd_cvm,
            anos=janela,
        )

        erros_dupont = (
            int(
                dupont["STATUS"]
                .astype(str)
                .eq("ERRO")
                .sum()
            )
            if not dupont.empty
            and "STATUS" in dupont.columns
            else 0
        )

        adicionar_resultado(
            resultados,
            cd_cvm,
            empresa,
            "DuPont",
            "OK" if erros_dupont == 0 else "FALHA",
            (
                "Nenhum resíduo fora da tolerância"
                if erros_dupont == 0
                else f"{erros_dupont} erro(s)"
            ),
            caso,
        )

        validacao = validar_empresa(
            cd_cvm=cd_cvm,
            anos=janela,
        )

        status_invalidos = set(
            validacao["STATUS"]
            .astype(str)
            .unique()
        ).difference(STATUS_VALIDOS)

        geral = status_geral(validacao)

        adicionar_resultado(
            resultados,
            cd_cvm,
            empresa,
            "Validação executável",
            (
                "OK"
                if not status_invalidos
                and geral in STATUS_VALIDOS
                else "FALHA"
            ),
            (
                f"status geral={geral}; "
                f"linhas={len(validacao)}; "
                f"status inválidos={sorted(status_invalidos)}"
            ),
            caso,
        )

        # Um BLOQUEIO/ALERTA por dados não é falha do software.
        # O objetivo aqui é confirmar que o sistema o detecta e propaga.
        contagem_status = (
            validacao["STATUS"]
            .value_counts()
            .to_dict()
        )

        adicionar_resultado(
            resultados,
            cd_cvm,
            empresa,
            "Tratamento de ressalvas",
            "OK",
            str(contagem_status),
            caso,
        )

        if profundo:
            relatorio = gerar_relatorio(
                cd_cvm=cd_cvm,
                nome_empresa=empresa,
                anos=janela,
            )

            ausentes_rel = (
                CHAVES_RELATORIO
                .difference(relatorio.keys())
                if isinstance(relatorio, dict)
                else CHAVES_RELATORIO
            )

            status_rel = (
                relatorio.get("STATUS_VALIDACAO")
                if isinstance(relatorio, dict)
                else None
            )

            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                "Relatório estrutural",
                (
                    "OK"
                    if not ausentes_rel
                    and status_rel == geral
                    else "FALHA"
                ),
                (
                    f"status relatório={status_rel}; "
                    f"status validação={geral}; "
                    f"chaves ausentes={sorted(ausentes_rel)}"
                ),
                caso,
            )

            cards = montar_cards_dashboard(
                indicadores,
                tuple(janela),
            )

            graficos = montar_graficos_dashboard(
                indicadores,
                tuple(janela),
            )

            cod_cards = (
                set(
                    cards["INDICADOR"]
                    .astype(str)
                    .tolist()
                )
                if isinstance(cards, pd.DataFrame)
                and not cards.empty
                and "INDICADOR" in cards.columns
                else set()
            )

            faltam_cards = set(
                INDICADORES_ESPERADOS
            ).difference(cod_cards)

            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                "Dashboard — cards",
                (
                    "OK"
                    if not faltam_cards
                    else "FALHA"
                ),
                (
                    f"cards={len(cod_cards)}; "
                    f"faltantes={sorted(faltam_cards)}"
                ),
                caso,
            )

            chaves_graficos = (
                set(graficos.keys())
                if isinstance(graficos, dict)
                else set()
            )

            faltam_graficos = (
                CHAVES_GRAFICOS
                .difference(chaves_graficos)
            )

            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                "Dashboard — gráficos",
                (
                    "OK"
                    if not faltam_graficos
                    else "FALHA"
                ),
                (
                    "gráficos esperados presentes"
                    if not faltam_graficos
                    else "faltantes="
                    + ", ".join(sorted(faltam_graficos))
                ),
                caso,
            )

            total_series = sum(
                len(
                    figura.data
                )
                for figura
                in graficos.values()
            )

            if layout_codigo == LAYOUT_PADRAO:
                series_ok = (
                    total_series == 12
                )
                detalhe_series = (
                    f"layout padrão; séries={total_series}; "
                    "esperado=12"
                )

            elif tratamento_setorial:
                rentabilidade_series = len(
                    graficos[
                        "rentabilidade"
                    ].data
                )

                outras_series = sum(
                    len(
                        graficos[
                            chave
                        ].data
                    )
                    for chave in [
                        "estrutura",
                        "liquidez",
                        "icj",
                        "ga",
                    ]
                )

                series_ok = (
                    total_series == 2
                    and rentabilidade_series == 2
                    and outras_series == 0
                )

                detalhe_series = (
                    f"layout={layout_codigo}; "
                    f"séries totais={total_series}; "
                    f"rentabilidade={rentabilidade_series}; "
                    f"demais={outras_series}"
                )

            else:
                series_ok = (
                    total_series == 0
                )

                detalhe_series = (
                    f"layout={layout_codigo}; "
                    f"séries totais={total_series}; "
                    "esperado=0"
                )

            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                "Dashboard — séries por layout",
                "OK" if series_ok else "FALHA",
                detalhe_series,
                caso,
            )

            linha_catalogo = pd.read_parquet(
                ARQUIVO_EMPRESAS
            )
            linha_catalogo["CD_CVM"] = (
                linha_catalogo["CD_CVM"]
                .astype("string")
                .str.strip()
                .str.zfill(6)
            )
            linha_empresa = linha_catalogo.loc[
                linha_catalogo["CD_CVM"].eq(cd_cvm)
            ]

            cnpj = (
                str(linha_empresa.iloc[0]["CNPJ_CIA"])
                if not linha_empresa.empty
                else ""
            )

            identificacao = {
                "DENOM_CIA": empresa,
                "CD_CVM": cd_cvm,
                "CNPJ_CIA": cnpj,
                "LAYOUT_CVM": (
                    layout_codigo
                    or "N/D"
                ),
                "TRATAMENTO_SETORIAL": (
                    tratamento_setorial
                ),
            }

            conteudo_xlsx = gerar_excel_sistema(
                identificacao=identificacao,
                bp_ativo=demonstracoes.get("BPA"),
                bp_passivo=demonstracoes.get("BPP"),
                dre=demonstracoes.get("DRE"),
                indicadores=indicadores,
                relatorio=relatorio,
                validacao=validacao,
                anos=janela,
                status_validacao=status_rel,
            )

            nome_xlsx = nome_arquivo_exportacao(
                identificacao,
                anos=janela,
            )

            abas = nomes_abas_xlsx(
                conteudo_xlsx
            )

            adicionar_resultado(
                resultados,
                cd_cvm,
                empresa,
                "Exportação XLSX",
                (
                    "OK"
                    if abas == ABAS_EXPORTACAO
                    and len(conteudo_xlsx) > 5000
                    and nome_xlsx.lower().endswith(".xlsx")
                    else "FALHA"
                ),
                (
                    f"{len(conteudo_xlsx):,} bytes; "
                    f"abas={abas}; nome={nome_xlsx}"
                ),
                caso,
            )

            if tratamento_setorial:
                wb = load_workbook(
                    BytesIO(
                        conteudo_xlsx
                    ),
                    data_only=False,
                    read_only=True,
                )

                ws_ind = wb[
                    "Indicadores"
                ]

                valores_ind = [
                    celula.value
                    for linha
                    in ws_ind.iter_rows()
                    for celula
                    in linha
                ]

                ws_dre = wb[
                    "DRE"
                ]

                valores_dre = [
                    celula.value
                    for linha
                    in ws_dre.iter_rows()
                    for celula
                    in linha
                ]

                ws_id = wb[
                    "Identificação"
                ]

                id_campos = {
                    str(
                        ws_id.cell(
                            linha,
                            1,
                        ).value
                        or ""
                    ): ws_id.cell(
                        linha,
                        2,
                    ).value
                    for linha
                    in range(
                        1,
                        ws_id.max_row
                        + 1,
                    )
                }

                export_setorial_ok = (
                    "N/A" in valores_ind
                    and "N/A" in valores_dre
                    and id_campos.get(
                        "Layout CVM"
                    )
                    == layout_codigo
                )

                adicionar_resultado(
                    resultados,
                    cd_cvm,
                    empresa,
                    "Exportação setorial — N/A e layout",
                    (
                        "OK"
                        if export_setorial_ok
                        else "FALHA"
                    ),
                    (
                        f"N/A indicadores={'N/A' in valores_ind}; "
                        f"N/A DRE={'N/A' in valores_dre}; "
                        f"layout exportado="
                        f"{id_campos.get('Layout CVM')}"
                    ),
                    caso,
                )

                wb.close()

    except Exception as exc:
        adicionar_resultado(
            resultados,
            cd_cvm,
            empresa,
            "Execução da companhia",
            "FALHA",
            (
                f"{type(exc).__name__}: {exc}\n"
                + traceback.format_exc(limit=6)
            ),
            caso,
        )

    duracao = time.perf_counter() - inicio

    adicionar_resultado(
        resultados,
        cd_cvm,
        empresa,
        "Tempo de execução",
        "INFO",
        f"{duracao:.2f} s",
        caso,
    )


# ------------------------------------------------------------
# EXECUÇÃO
# ------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description="Auditoria adversarial do Sistema CVM."
    )

    parser.add_argument(
        "--modo",
        choices=["rapido", "completo"],
        default="rapido",
        help="rapido = amostra adversarial; completo = núcleo em todas as empresas",
    )

    args = parser.parse_args()

    inicio_total = time.perf_counter()
    resultados: list[dict] = []

    print("=" * 78)
    print("AUDITORIA ADVERSARIAL — SISTEMA CVM")
    print("=" * 78)
    print(f"Modo: {args.modo.upper()}")
    print()

    catalogo = carregar_catalogo()

    print(f"Empresas no catálogo: {len(catalogo):,}")

    homonimos = auditar_catalogo(
        catalogo,
        resultados,
    )

    casos = descobrir_casos(
        catalogo,
        homonimos,
    )

    caminho_casos = (
        DATA_DIR
        / "auditoria_adversarial_casos.csv"
    )

    casos.to_csv(
        caminho_casos,
        index=False,
        sep=";",
        encoding="utf-8-sig",
    )

    print(
        f"Amostra adversarial profunda: "
        f"{len(casos):,} companhia(s)"
    )

    if not casos.empty:
        print()
        print("CASOS SELECIONADOS")
        print("-" * 78)

        for _, linha in casos.iterrows():
            print(
                f"{linha['CD_CVM']} | "
                f"{linha['EMPRESA']} | "
                f"{linha['CASO']}"
            )

    print()
    print("TESTES PROFUNDOS")
    print("-" * 78)

    for pos, (_, linha) in enumerate(
        casos.iterrows(),
        start=1,
    ):
        print(
            f"[{pos:02d}/{len(casos):02d}] "
            f"{linha['CD_CVM']} | "
            f"{linha['EMPRESA']}"
        )

        testar_empresa(
            cd_cvm=linha["CD_CVM"],
            empresa=linha["EMPRESA"],
            caso=linha["CASO"],
            resultados=resultados,
            profundo=True,
        )

    if args.modo == "completo":
        ja_testados = set(
            casos["CD_CVM"].astype(str)
        )

        restantes = catalogo.loc[
            ~catalogo["CD_CVM"].isin(
                ja_testados
            )
        ].reset_index(drop=True)

        print()
        print("VARREDURA COMPLETA DO NÚCLEO")
        print("-" * 78)
        print(
            f"Companhias adicionais: "
            f"{len(restantes):,}"
        )

        for pos, (_, linha) in enumerate(
            restantes.iterrows(),
            start=1,
        ):
            if (
                pos == 1
                or pos % 25 == 0
                or pos == len(restantes)
            ):
                print(
                    f"[{pos:03d}/{len(restantes):03d}] "
                    f"{linha['CD_CVM']} | "
                    f"{linha['DENOM_CIA']}"
                )

            testar_empresa(
                cd_cvm=linha["CD_CVM"],
                empresa=str(
                    linha["DENOM_CIA"]
                ),
                caso="varredura completa",
                resultados=resultados,
                profundo=False,
            )

    df_resultados = pd.DataFrame(
        resultados
    )

    caminho_resultados = (
        DATA_DIR
        / "auditoria_adversarial_resultados.csv"
    )

    df_resultados.to_csv(
        caminho_resultados,
        index=False,
        sep=";",
        encoding="utf-8-sig",
    )

    falhas = df_resultados.loc[
        df_resultados["STATUS"].eq("FALHA")
    ]

    oks = int(
        df_resultados["STATUS"]
        .eq("OK")
        .sum()
    )

    infos = int(
        df_resultados["STATUS"]
        .eq("INFO")
        .sum()
    )

    duracao_total = (
        time.perf_counter()
        - inicio_total
    )

    print()
    print("=" * 78)
    print("RESULTADO FINAL")
    print("=" * 78)
    print(f"Testes OK:       {oks:,}")
    print(f"Informativos:    {infos:,}")
    print(f"Falhas técnicas: {len(falhas):,}")
    print(f"Tempo total:     {duracao_total:.1f} s")
    print()
    print(f"Resultados: {caminho_resultados}")
    print(f"Casos:      {caminho_casos}")

    if falhas.empty:
        print()
        print("STATUS DA AUDITORIA: APROVADO")
        print(
            "Nenhuma falha técnica foi encontrada "
            "nos testes executados."
        )
        return 0

    print()
    print("STATUS DA AUDITORIA: REPROVADO")
    print()
    print("PRIMEIRAS FALHAS")
    print("-" * 78)

    colunas = [
        "CD_CVM",
        "EMPRESA",
        "CASO",
        "TESTE",
        "DETALHE",
    ]

    print(
        falhas[colunas]
        .head(20)
        .to_string(index=False)
    )

    return 1


if __name__ == "__main__":
    raise SystemExit(main())
